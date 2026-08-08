import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def hex_to_rgb(hex_str):
    hex_str = hex_str.lstrip('#')
    return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))

BG_COLOR = hex_to_rgb('0F1113')
TEXT_COLOR = RGBColor(240, 240, 240) # Off-white for readability
TERRACOTTA = hex_to_rgb('B8632A')
GOLD = hex_to_rgb('C9973F')

HEADING_FONT = 'Georgia'
BODY_FONT = 'Calibri'

def apply_styling(shape, font_name, font_size, color, bold=False, italic=False, align=PP_ALIGN.LEFT):
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    for paragraph in tf.paragraphs:
        paragraph.alignment = align
        for run in paragraph.runs:
            run.font.name = font_name
            run.font.size = Pt(font_size)
            run.font.color.rgb = color
            run.font.bold = bold
            run.font.italic = italic

def add_background_and_accents(slide):
    # Set background color
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

    # Add accent bar (Top)
    # Using two thin rectangles for terracotta and gold
    left1 = Inches(0)
    top1 = Inches(0)
    width1 = Inches(5)
    height1 = Inches(0.15)
    rect1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left1, top1, width1, height1)
    rect1.fill.solid()
    rect1.fill.fore_color.rgb = TERRACOTTA
    rect1.line.fill.background()

    left2 = Inches(5)
    top2 = Inches(0)
    width2 = Inches(5)
    height2 = Inches(0.15)
    rect2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left2, top2, width2, height2)
    rect2.fill.solid()
    rect2.fill.fore_color.rgb = GOLD
    rect2.line.fill.background()

def create_title_slide(prs, title_text, subtitle_text):
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    add_background_and_accents(slide)
    
    # Title
    left = Inches(1)
    top = Inches(2.5)
    width = Inches(8)
    height = Inches(1.5)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.add_paragraph()
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = HEADING_FONT
    p.font.size = Pt(44)
    p.font.color.rgb = TERRACOTTA
    p.font.bold = True
    
    # Subtitle
    left = Inches(1)
    top = Inches(4.5)
    width = Inches(8)
    height = Inches(1)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.add_paragraph()
    p.text = subtitle_text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = BODY_FONT
    p.font.size = Pt(24)
    p.font.color.rgb = GOLD

def create_content_slide(prs, title_text, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background_and_accents(slide)
    
    # Title
    left = Inches(0.8)
    top = Inches(0.8)
    width = Inches(8.4)
    height = Inches(1)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.add_paragraph()
    p.text = title_text
    p.font.name = HEADING_FONT
    p.font.size = Pt(36)
    p.font.color.rgb = GOLD
    p.font.bold = True
    
    # Content
    left = Inches(0.8)
    top = Inches(2.2)
    width = Inches(8.4)
    height = Inches(4.5)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.space_before = Pt(20)
        p.font.name = BODY_FONT
        p.font.size = Pt(24)
        p.font.color.rgb = TEXT_COLOR

def main():
    prs = Presentation()
    
    # Slide 1
    create_title_slide(prs, "Indian Folk Art Intelligence System", "Style Classification, Authenticity Detection & Cultural Heritage Preservation")
    
    # Slide 2
    create_content_slide(prs, "The Problem", [
        "Most people can't distinguish folk art styles like Warli vs Madhubani.",
        "Machine-printed replicas are sold online as genuine hand-painted originals, with no way for buyers to verify."
    ])
    
    # Slide 3
    create_content_slide(prs, "What This Project Does", [
        "One photo upload returns four things:",
        "► Style",
        "► Authenticity",
        "► Visual Explanation",
        "► Cultural Story"
    ])
    
    # Slide 4
    create_content_slide(prs, "System Overview", [
        "Five components working together:",
        "[ Style Classifier ] → [ Authenticity Detector ] → [ Explainability (Grad-CAM) ] → [ Similarity Search ] → [ Cultural Context Engine ]"
    ])
    
    # Slide 5
    create_content_slide(prs, "Style Classifier", [
        "Architecture: ResNet18 CNN (transfer learning)",
        "Coverage: 8 distinct folk art styles",
        "Performance: 90% validation accuracy",
        "Note: Pretrained on ImageNet, fine-tuned on custom folk art dataset."
    ])
    
    # Slide 6
    create_content_slide(prs, "Authenticity Detector", [
        "Attempt 1: CNN approach, ~55-60% accuracy (failed).",
        "Attempt 2: Random Forest on engineered features (sharpness, saturation, edge density, noise), 91% accuracy (succeeded).",
        "Takeaway: The right tool matters more than the fanciest tool."
    ])
    
    # Slide 7
    create_content_slide(prs, "Explainability (Grad-CAM)", [
        "Shows which part of the painting the AI focused on to make its decision.",
        "Builds trust, ensures the model is not a black box.",
        "Highlights stylistic motifs critical to classification."
    ])
    
    # Slide 8
    create_content_slide(prs, "Similarity Search", [
        "Finds the most visually similar paintings from the dataset using AI embeddings.",
        "Acts as a \"you might also like\" discovery engine for folk art."
    ])
    
    # Slide 9
    create_content_slide(prs, "Cultural Context Engine", [
        "Provides real history, regional origins, and symbolism per style.",
        "Grounds technical AI results in cultural heritage."
    ])
    
    # Slide 10
    create_content_slide(prs, "Results", [
        "► 90% Style Classification Accuracy",
        "► 91% Authenticity Detection Accuracy",
        "Evaluated on validation data the model never trained on."
    ])
    
    # Slide 11
    create_content_slide(prs, "Tech Stack", [
        "Models: PyTorch",
        "Backend: FastAPI",
        "Frontend: React",
        "A full working web app, not just a notebook."
    ])
    
    # Slide 12
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background_and_accents(slide)
    
    left = Inches(1)
    top = Inches(3)
    width = Inches(8)
    height = Inches(1.5)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.add_paragraph()
    p.text = "AI that explains itself, protects buyers, and teaches culture — all in one tool."
    p.alignment = PP_ALIGN.CENTER
    p.font.name = HEADING_FONT
    p.font.size = Pt(36)
    p.font.color.rgb = TERRACOTTA
    p.font.bold = True
    
    left = Inches(1)
    top = Inches(5.5)
    width = Inches(8)
    height = Inches(1)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.add_paragraph()
    p.text = "Thank you / Questions"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = BODY_FONT
    p.font.size = Pt(24)
    p.font.color.rgb = GOLD

    output_path = "Indian_Folk_Art_Intelligence_System.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")
    print(f"Total slides: {len(prs.slides)}")

if __name__ == '__main__':
    main()
